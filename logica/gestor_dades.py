"""
Gestor de dades: llegir/escriure el fitxer Excel local.
"""

import os
from datetime import datetime
from openpyxl import Workbook, load_workbook


# ============================================================
#  CONFIGURACIÓ
# ============================================================

CARPETA_TREBALL = r"C:\CutterPython\dades"
FITXER_EXCEL = os.path.join(CARPETA_TREBALL, "cutter.xlsx")
CARPETA_D2K = r"C:\MacrosTall"


# ============================================================
#  CAPÇALERA DE L'EXCEL
# ============================================================

CAPCALERA = [
    "Mida", "Angle Esq", "Angle Dret", "Quantitat",
    "Perfil", "Client", "Gruix", "Descripció",
    "Número", "Tipus Obertura", "Color", "Auxiliar", "Codi Feina"
]

FILA_INICI = 6


# ============================================================
#  CREAR EL FITXER EXCEL SI NO EXISTEIX
# ============================================================

def crear_fitxer_si_no_existeix():
    if not os.path.exists(CARPETA_TREBALL):
        os.makedirs(CARPETA_TREBALL)

    if not os.path.exists(FITXER_EXCEL):
        wb = Workbook()
        ws = wb.active
        ws.title = "Talls"

        ws["A1"] = "Tanqui-Tanqui, s.l."
        ws["A2"] = "LLISTA DE TALLS"

        for i, titol in enumerate(CAPCALERA, start=1):
            ws.cell(row=5, column=i, value=titol)

        amplades = [12, 10, 10, 10, 10, 12, 8, 18, 8, 25, 10, 10, 12]
        for i, amp in enumerate(amplades, start=1):
            ws.column_dimensions[chr(64 + i)].width = amp

        wb.save(FITXER_EXCEL)


# ============================================================
#  LLEGIR TOTS ELS TALLS
# ============================================================

def llegir_talls():
    crear_fitxer_si_no_existeix()

    wb = load_workbook(FITXER_EXCEL)
    ws = wb.active

    talls = []
    fila = FILA_INICI

    while ws.cell(row=fila, column=1).value is not None:
        tall = {
            "mida": ws.cell(row=fila, column=1).value,
            "angle_esq": ws.cell(row=fila, column=2).value,
            "angle_dret": ws.cell(row=fila, column=3).value,
            "quantitat": ws.cell(row=fila, column=4).value,
            "perfil": ws.cell(row=fila, column=5).value,
            "client": ws.cell(row=fila, column=6).value,
            "gruix": ws.cell(row=fila, column=7).value,
            "descripcio": ws.cell(row=fila, column=8).value,
            "numero": ws.cell(row=fila, column=9).value,
            "tipus_obertura": ws.cell(row=fila, column=10).value,
            "color": ws.cell(row=fila, column=11).value,
            "ordre_original": ws.cell(row=fila, column=12).value,
            "codi_feina": ws.cell(row=fila, column=13).value,
            "fila": fila,
        }
        talls.append(tall)
        fila += 1

    wb.close()
    return talls


# ============================================================
#  AFEGIR UN TALL
# ============================================================

def afegir_tall(mida, angle_esq, angle_dret, quantitat, perfil,
                client, gruix, descripcio, tipus_obertura, color,
                codi_feina):
    crear_fitxer_si_no_existeix()

    wb = load_workbook(FITXER_EXCEL)
    ws = wb.active

    fila = FILA_INICI
    while ws.cell(row=fila, column=1).value is not None:
        fila += 1

    ws.cell(row=fila, column=1, value=mida)
    ws.cell(row=fila, column=2, value=angle_esq)
    ws.cell(row=fila, column=3, value=angle_dret)
    ws.cell(row=fila, column=4, value=quantitat)
    ws.cell(row=fila, column=5, value=perfil)
    ws.cell(row=fila, column=6, value=client)
    ws.cell(row=fila, column=7, value=gruix)
   # Evitar duplicar el prefix del client
    if descripcio.startswith(f"{client[:3]}-"):
        desc_final = descripcio
    else:
        desc_final = f"{client[:3]}-{descripcio}"

    ws.cell(row=fila, column=8, value=desc_final)
    ws.cell(row=fila, column=9, value=fila)
    ws.cell(row=fila, column=10, value=tipus_obertura)
    ws.cell(row=fila, column=11, value=color)
    ws.cell(row=fila, column=12, value=fila)  # Ordre original
    ws.cell(row=fila, column=13, value=codi_feina)
    
    wb.save(FITXER_EXCEL)
    wb.close()

    return fila


# ============================================================
#  AFEGIR MÚLTIPLES TALLS
# ============================================================

def afegir_talls(llista_talls, client, tipus_obertura, color, codi_feina):
    crear_fitxer_si_no_existeix()

    wb = load_workbook(FITXER_EXCEL)
    ws = wb.active

    fila = FILA_INICI
    while ws.cell(row=fila, column=1).value is not None:
        fila += 1

    comptador = 0
    for tall in llista_talls:
        ws.cell(row=fila, column=1, value=tall["mida"])
        ws.cell(row=fila, column=2, value=tall["angle_esq"])
        ws.cell(row=fila, column=3, value=tall["angle_dret"])
        ws.cell(row=fila, column=4, value=tall["quantitat"])
        ws.cell(row=fila, column=5, value=tall["perfil"])
        ws.cell(row=fila, column=6, value=client)
        ws.cell(row=fila, column=7, value=tall["gruix"])
                # Evitar duplicar el prefix
        desc_tall = tall['descripcio']
        if desc_tall.startswith(f"{client[:3]}-"):
            desc_final = desc_tall
        else:
            desc_final = f"{client[:3]}-{desc_tall}"

        ws.cell(row=fila, column=8, value=desc_final)
        ws.cell(row=fila, column=9, value=fila)
        ws.cell(row=fila, column=10, value=tipus_obertura)
        ws.cell(row=fila, column=11, value=color)
        ws.cell(row=fila, column=12, value=fila)  # Ordre original
        ws.cell(row=fila, column=13, value=codi_feina)
        fila += 1
        comptador += 1

    wb.save(FITXER_EXCEL)
    wb.close()

    return comptador


# ============================================================
#  ESBORRAR TOTS ELS TALLS
# ============================================================

def esborrar_talls():
    crear_fitxer_si_no_existeix()

    wb = load_workbook(FITXER_EXCEL)
    ws = wb.active

    ws.delete_rows(FILA_INICI, ws.max_row)

    wb.save(FITXER_EXCEL)
    wb.close()


# ============================================================
#  COMPTAR TALLS
# ============================================================

def comptar_talls():
    talls = llegir_talls()
    return len(talls)


# ============================================================
#  ORDENAR
# ============================================================

def ordenar_per_tipus(talls):
    """
    Ordena: Marcs primer, després Fulles, després Goterons, Rivets, Inversor/Creuament, etc.
    Dins de cada tipus, ordena per PERFIL.
    """
    def clau(t):
        desc = (t.get('descripcio') or '').upper()
        perfil = (t.get('perfil') or '').upper()

        if 'MARC' in desc:
            tipus = 0
        elif 'FULLA' in desc:
            tipus = 1
        elif 'GOTERON' in desc:
            tipus = 2
        elif 'RIVET' in desc:
            tipus = 3
        elif 'INVERSOR' in desc or 'CREUAMENT' in desc:
            tipus = 4
        else:
            tipus = 5

        return (tipus, perfil)

    return sorted(talls, key=clau)


def ordenar_per_mida_desc(talls):
    """
    Ordena per tipus, perfil i mida (de major a menor).
    """
    def clau(t):
        desc = (t.get('descripcio') or '').upper()
        perfil = (t.get('perfil') or '').upper()
        mida = t.get('mida') or 0

        if 'MARC' in desc:
            tipus = 0
        elif 'FULLA' in desc:
            tipus = 1
        elif 'GOTERON' in desc:
            tipus = 2
        elif 'RIVET' in desc:
            tipus = 3
        elif 'INVERSOR' in desc or 'CREUAMENT' in desc:
            tipus = 4
        else:
            tipus = 5

        return (tipus, perfil, -mida)

    return sorted(talls, key=clau)


def ordenar_per_codi_feina(talls):
    """Ordena els talls pel codi de feina."""
    return sorted(talls, key=lambda t: (t.get('codi_feina') or ''))


# ============================================================
#  EXPORTAR A D2K
# ============================================================

def exportar_d2k(codi_barres_inicial="12345678901"):
    talls = llegir_talls()

    if not talls:
        return None

    if not os.path.exists(CARPETA_D2K):
        print(f"ERROR: No existeix la carpeta {CARPETA_D2K}")
        return None

    nom_fitxer = f"ORDRE_TALL_D2K_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"
    ruta_fitxer = os.path.join(CARPETA_D2K, nom_fitxer)

    try:
        codi_barres = float(codi_barres_inicial)
    except (ValueError, TypeError):
        codi_barres = 12345678901

    try:
        with open(ruta_fitxer, "w", encoding="utf-8") as f:
            perfil_anterior = ""
            color_anterior = ""

            for tall in talls:
                perfil_actual = tall["perfil"]
                color_actual = tall["color"] or "-"

                if perfil_actual != perfil_anterior or color_actual != color_anterior:
                    gruix = tall["gruix"] or 50
                    f.write(f"B,{perfil_actual},{color_actual},1,{gruix:.0f}\n")
                    perfil_anterior = perfil_actual
                    color_anterior = color_actual

                mida_d2k = int(tall["mida"] * 10) if tall["mida"] else 0
                ae_d2k = int(tall["angle_esq"] * 10) if tall["angle_esq"] else 450
                ad_d2k = int(tall["angle_dret"] * 10) if tall["angle_dret"] else 450
                quantitat = int(tall["quantitat"]) if tall["quantitat"] else 1

                tipus_peca = tall["descripcio"] or "Upright"
                client = tall["client"] or "CLIENT"
                codi_feina = tall["codi_feina"] or ""

                f.write(
                    f"P,{mida_d2k:06d},0,{ae_d2k:04d},{ad_d2k:04d},{quantitat:02d},,,"
                    f"{tipus_peca},,{client} {codi_feina},,,,"
                    f"{int(codi_barres):013d}\n"
                )

                codi_barres += 1

        return ruta_fitxer

    except Exception as e:
        print(f"ERROR escrivint el fitxer D2K: {e}")
        return None
    
def importar_d2k(ruta_fitxer):
    """
    Llegeix un fitxer D2K i retorna una llista de talls.
    Equivalent al RecuperarD2K del VBA.
    
    Args:
        ruta_fitxer: Ruta del fitxer .txt D2K
    
    Returns:
        list: Llista de diccionaris amb els talls llegits.
    """
    talls = []
    perfil_actual = ""
    color_actual = "BLANC"
    gruix_actual = 50
    
    try:
        with open(ruta_fitxer, "r", encoding="utf-8", errors="ignore") as f:
            for linia in f:
                linia = linia.strip()
                if not linia:
                    continue
                
                parts = linia.split(",")
                
                # ---- Línia B (canvi de perfil/color) ----
                if len(parts) >= 5 and parts[0] == "B":
                    perfil_actual = parts[1] if len(parts) > 1 else ""
                    color_actual = parts[2] if len(parts) > 2 and parts[2] else "BLANC"
                    if len(parts) > 4:
                        try:
                            gruix_actual = float(parts[4])
                        except (ValueError, TypeError):
                            pass
                
                # ---- Línia P (peça) ----
                elif len(parts) >= 12 and parts[0] == "P":
                    try:
                        mida = float(parts[1]) / 10 if parts[1] else 0
                    except (ValueError, TypeError):
                        mida = 0
                    
                    try:
                        ae = float(parts[3]) / 10 if parts[3] else 45
                    except (ValueError, TypeError):
                        ae = 45
                    
                    try:
                        ad = float(parts[4]) / 10 if parts[4] else 45
                    except (ValueError, TypeError):
                        ad = 45
                    
                    try:
                        quantitat = int(parts[5]) if parts[5] else 1
                    except (ValueError, TypeError):
                        quantitat = 1
                    
                    desc = parts[9] if len(parts) > 9 else ""
                    
                    # Extreure client i codi feina
                    client_i_codi = parts[11].strip() if len(parts) > 11 else ""
                    parts_client = client_i_codi.split(" ")
                    
                    if len(parts_client) >= 2:
                        codi_final = parts_client[-1]
                        client_final = client_i_codi[:len(client_i_codi) - len(codi_final) - 1]
                    else:
                        client_final = client_i_codi
                        codi_final = ""
                    
                    tall = {
                        "mida": mida,
                        "angle_esq": ae,
                        "angle_dret": ad,
                        "quantitat": quantitat,
                        "perfil": perfil_actual,
                        "client": client_final,
                        "gruix": gruix_actual,
                        "descripcio": desc,
                        "numero": len(talls) + 1,
                        "tipus_obertura": "",
                        "color": color_actual,
                        "codi_feina": codi_final,
                        "ordre_original": len(talls) + 1,
                    }
                    talls.append(tall)
        
        return talls
    
    except Exception as e:
        print(f"Error llegint el fitxer D2K: {e}")
        return []